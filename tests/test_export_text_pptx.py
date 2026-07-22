from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "export_text_pptx.py"
SPEC = importlib.util.spec_from_file_location("export_text_pptx", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(MODULE)


class WhiteClassroomPptxExporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "ppt-text-structured.json"
        self.approval = self.root / "approval.json"
        self.output = self.root / "exports"
        self.payload = self._payload()
        self._write_inputs()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _payload(self) -> dict:
        return {
            "schema_version": "1.0",
            "record_type": "ppt-text",
            "lesson_id": "sample-lesson",
            "status": "draft",
            "generated_at": "2026-07-19T08:00:00Z",
            "text_contract_sha256": "1" * 64,
            "ppt_text_sha256": "2" * 64,
            "slides": [
                {
                    "slide_id": "SLIDE-01",
                    "page": 1,
                    "title": "示例课文·第二课时",
                    "screen_text": "示例文本的详略\n\n自创示例材料｜课堂练习",
                    "teacher_notes": "开课前上屏，只读劝降一段。",
                    "task_id": "TASK-01",
                    "objective_ids": ["OBJ-01"],
                    "script_ids": ["SCRIPT-01"],
                    "answer_ids": [],
                    "reveal_state": "neutral",
                },
                {
                    "slide_id": "SLIDE-02",
                    "page": 2,
                    "title": "今天只读一段",
                    "screen_text": (
                        "上节课的脉络：出使 → 被扣 → 劝降\n\n"
                        "核心问题：作者为什么这样安排详略？\n\n"
                        "活动：①默读卫律劝降段，圈画详写处；②同桌比较；"
                        "③用一句话说明详略安排的作用。\n"
                        "时间：5 分钟\n"
                        "成果：一句结论和两处文本依据\n"
                        "交流：同桌互证后，由一人向全班汇报。"
                    ),
                    "teacher_notes": "抛出核心问题后齐读卫律段。",
                    "task_id": "TASK-01",
                    "objective_ids": ["OBJ-01"],
                    "script_ids": ["SCRIPT-01"],
                    "answer_ids": [],
                    "reveal_state": "prompt",
                },
            ],
        }

    def _write_inputs(self) -> None:
        self.source.write_text(
            json.dumps(self.payload, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        source_bytes = self.source.read_bytes()
        source_hash = hashlib.sha256(source_bytes).hexdigest()
        approval = {
            "approval_id": "r0021-ppt-text",
            "approved_at": "2026-07-19T08:05:00Z",
            "confirmation": {
                "path": "confirmation.json",
                "sha256": "3" * 64,
                "size": 1,
            },
            "expected_revision": 20,
            "inputs": [
                {
                    "path": self.source.name,
                    "role": "ppt-structured",
                    "sha256": source_hash,
                    "size": len(source_bytes),
                }
            ],
            "lesson_id": self.payload["lesson_id"],
            "next_stage": "release",
            "record_type": "user-stage-approval",
            "resulting_revision": 21,
            "schema_version": "1.0",
            "stage": "ppt-text",
            "stage_index": 9,
            "user_id": "teacher-test",
        }
        self.approval.write_text(
            json.dumps(approval, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

    def export(self, output: Path | None = None) -> dict:
        return MODULE.export_text_pptx(
            self.source,
            output or self.output,
            approval_record=self.approval,
            name="示例课文白底课件",
            font_checker=lambda _font: True,
            created_at=datetime(2026, 7, 19, 8, 10, tzinfo=timezone.utc),
        )

    def test_exports_editable_white_classroom_presentation_with_notes_and_lineage(self) -> None:
        report = self.export()
        bundle = self.output / "示例课文白底课件.export"
        pptx_path = bundle / "示例课文白底课件.pptx"
        report_path = bundle / "示例课文白底课件.export-report.json"
        self.assertTrue(pptx_path.is_file())
        self.assertTrue(report_path.is_file())
        self.assertEqual(report["schema_version"], "yuwen-courseware.text-pptx-export/v1")
        self.assertEqual(report["source"]["sha256"], hashlib.sha256(self.source.read_bytes()).hexdigest())
        self.assertEqual(report["artifact"]["sha256"], hashlib.sha256(pptx_path.read_bytes()).hexdigest())
        self.assertEqual(report["artifact"]["slide_count"], 2)

        presentation = Presentation(pptx_path)
        self.assertEqual(len(presentation.slides), 2)
        self.assertEqual(presentation.slide_width, 12192000)
        self.assertEqual(presentation.slide_height, 6858000)
        self.assertEqual(presentation.core_properties.subject, "高中语文白底课堂课件")
        self.assertEqual(presentation.core_properties.author, "语文优课")
        self.assertEqual(presentation.core_properties.last_modified_by, "语文优课")
        self.assertEqual(
            report["lineage"]["layout_policy"],
            "16:9-white-black-no-autoshrink",
        )
        for index, slide in enumerate(presentation.slides, 1):
            self.assertTrue(all(shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX for shape in slide.shapes))
            self.assertEqual(slide.background.fill.fore_color.rgb, RGBColor(255, 255, 255))
            self.assertIn(self.payload["slides"][index - 1]["teacher_notes"], slide.notes_slide.notes_text_frame.text)
            visible_runs = [
                run
                for shape in slide.shapes
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.text
            ]
            self.assertTrue(visible_runs)
            self.assertTrue(all(run.font.name == "Source Han Serif SC" for run in visible_runs))
            self.assertTrue(all(run.font.size is not None and run.font.size.pt >= 28 for run in visible_runs))
            self.assertTrue(all(run.font.color.rgb == RGBColor(0, 0, 0) for run in visible_runs))
        first_title = presentation.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        second_title = presentation.slides[1].shapes[0].text_frame.paragraphs[0].runs[0]
        self.assertGreaterEqual(first_title.font.size.pt, 72)
        self.assertEqual(second_title.font.size.pt, 40)
        second_page_text = "\n".join(shape.text for shape in presentation.slides[1].shapes)
        for label in ("核心问题：", "活动：", "时间：", "成果：", "交流："):
            with self.subTest(label=label):
                self.assertIn(label, second_page_text)

        editable_copy = self.root / "可编辑检查.pptx"
        first_title.text = "可编辑标题"
        presentation.save(editable_copy)
        self.assertEqual(
            Presentation(editable_copy).slides[0].shapes[0].text_frame.paragraphs[0].text,
            "可编辑标题",
        )

    def test_same_inputs_create_byte_identical_pptx(self) -> None:
        first = self.export(self.root / "exports-one")
        second = self.export(self.root / "exports-two")
        first_path = self.root / "exports-one" / "示例课文白底课件.export" / "示例课文白底课件.pptx"
        second_path = self.root / "exports-two" / "示例课文白底课件.export" / "示例课文白底课件.pptx"
        self.assertEqual(first["artifact"]["sha256"], second["artifact"]["sha256"])
        self.assertEqual(first_path.read_bytes(), second_path.read_bytes())

    def test_overflow_fails_in_chinese_without_creating_bundle(self) -> None:
        self.payload["slides"] = [self.payload["slides"][0]]
        self.payload["slides"][0]["screen_text"] = (
            "问题：分析这段文字。\n活动：圈画并讨论。\n时间：5 分钟\n"
            "成果：一份分析结论\n交流：全班汇报\n"
            + "\n".join("超长屏显内容" * 8 for _ in range(40))
        )
        self._write_inputs()
        with self.assertRaisesRegex(MODULE.ExportError, r"第 1 页.*超出.*行"):
            self.export()
        self.assertFalse((self.output / "示例课文白底课件.export").exists())

    def test_every_interaction_component_must_be_visible_to_students(self) -> None:
        components = {
            "问题": "问题：作者为什么这样安排详略？",
            "活动": "活动：圈画详写处并交流作用。",
            "时间": "时间：5 分钟",
            "成果": "成果：一句结论和两处依据",
            "交流或提交方式": "交流：同桌互证后全班汇报。",
        }
        for missing_label, missing_line in components.items():
            with self.subTest(missing_label=missing_label):
                self.payload = self._payload()
                self.payload["slides"][1]["screen_text"] = "\n".join(
                    line for label, line in components.items() if label != missing_label
                )
                self.payload["slides"][1]["teacher_notes"] = missing_line
                self._write_inputs()
                with self.assertRaisesRegex(MODULE.ExportError, missing_label):
                    self.export()

    def test_neutral_task_still_requires_complete_visible_interaction(self) -> None:
        self.payload["slides"][1]["reveal_state"] = "neutral"
        self.payload["slides"][1]["screen_text"] = self.payload["slides"][1][
            "screen_text"
        ].replace("时间：5 分钟\n", "")
        self._write_inputs()

        with self.assertRaisesRegex(MODULE.ExportError, "时间"):
            self.export()

    def test_interaction_components_can_span_pages_of_the_same_task(self) -> None:
        self.payload["slides"][1]["screen_text"] = (
            "核心问题：作者为什么这样安排详略？\n"
            "时间：5 分钟\n成果：一句结论和两处依据"
        )
        self.payload["slides"].append(
            {
                "slide_id": "SLIDE-03",
                "page": 3,
                "title": "圈画与交流",
                "screen_text": (
                    "活动：①圈画详写处；②同桌比较；③修订结论。\n"
                    "交流：同桌互证后全班汇报。"
                ),
                "teacher_notes": "交流后回到核心问题。",
                "task_id": "TASK-01",
                "objective_ids": ["OBJ-01"],
                "script_ids": ["SCRIPT-01"],
                "answer_ids": [],
                "reveal_state": "neutral",
            }
        )
        self._write_inputs()

        report = self.export()

        self.assertEqual(report["artifact"]["slide_count"], 3)

    def test_interaction_content_can_follow_labels_on_the_same_or_next_line(self) -> None:
        self.payload["slides"][1]["screen_text"] = (
            "问题：\n作者为什么这样安排详略？\n"
            "活动：圈画详写处并交流作用。\n"
            "时间：\n5 分钟\n"
            "成果：一句结论和两处依据\n"
            "提交方式：\n在便笺上写下结论后交给同桌互评。"
        )
        self._write_inputs()

        report = self.export()

        self.assertEqual(report["artifact"]["slide_count"], 2)

    def test_empty_or_placeholder_interaction_values_fail_closed(self) -> None:
        cases = {
            "五项都为空": "问题：\n活动：\n时间：\n成果：\n交流：",
            "问题仍是占位语": (
                "问题：待填写\n"
                "活动：圈画详写处并交流作用。\n"
                "时间：5 分钟\n"
                "成果：一句结论和两处依据\n"
                "交流：同桌互证后全班汇报。"
            ),
        }
        for name, screen_text in cases.items():
            with self.subTest(name=name):
                self.payload = self._payload()
                self.payload["slides"][1]["screen_text"] = screen_text
                self._write_inputs()
                with self.assertRaisesRegex(MODULE.ExportError, "实际内容"):
                    self.export()

    def test_second_of_two_tasks_cannot_omit_an_interaction_component(self) -> None:
        self.payload["slides"].append(
            {
                "slide_id": "SLIDE-03",
                "page": 3,
                "title": "回看作者的选择",
                "screen_text": (
                    "问题：删去详写部分会怎样？\n"
                    "活动：小组改写并比较表达效果。\n"
                    "时间：4 分钟\n"
                    "成果：一份改写稿和比较结论"
                ),
                "teacher_notes": "第二个任务结束后回扣传记写法。",
                "task_id": "TASK-02",
                "objective_ids": ["OBJ-01"],
                "script_ids": ["SCRIPT-01"],
                "answer_ids": [],
                "reveal_state": "neutral",
            }
        )
        self._write_inputs()

        with self.assertRaisesRegex(MODULE.ExportError, "交流或提交方式"):
            self.export()

    def test_teacher_visible_fields_reject_backend_language_ids_statuses_and_emoji(self) -> None:
        cases = [
            *[("title", term) for term in MODULE.BACKEND_TERMS],
            *[
                ("screen_text", f"课堂说明 {prefix}-01")
                for prefix in ("OBJ", "TASK", "ASM", "ANS", "SCRIPT", "SLIDE", "SRC", "CLM", "MAT")
            ],
            ("teacher_notes", "当前状态 approved"),
            ("teacher_notes", "当前状态 student-blank"),
            ("teacher_notes", "status: review"),
            ("teacher_notes", "当前阶段 text-contract"),
            ("title", "课堂任务 ✨"),
        ]
        for field, value in cases:
            with self.subTest(field=field, value=value):
                self.payload = self._payload()
                self.payload["slides"][0][field] = value
                self._write_inputs()
                with self.assertRaisesRegex(MODULE.ExportError, "只保留课堂教学需要"):
                    self.export()

    def test_teacher_visible_fields_reject_black_box_production_and_check_notes(self) -> None:
        cases = [
            ("title", "schema_version：1.0"),
            ("screen_text", "JSON 校验报告：全部通过"),
            ("teacher_notes", "SHA-256：0123456789abcdef"),
            ("title", "manifest.json"),
            ("screen_text", "approval record 已写入"),
            ("teacher_notes", "检查结论：页面没有异常"),
            ("teacher_notes", "当前检查结论：页面没有异常"),
            ("teacher_notes", "课件制作过程：先运行脚本再导出文件"),
            ("teacher_notes", "本课件由 Python_PPTX 生成"),
            ("teacher_notes", "软件兼容性提醒：请用 WPS 打开"),
        ]
        for field, value in cases:
            with self.subTest(field=field, value=value):
                self.payload = self._payload()
                self.payload["slides"][0][field] = value
                self._write_inputs()
                with self.assertRaisesRegex(MODULE.ExportError, "只保留课堂教学需要"):
                    self.export()

    def test_normal_chinese_teaching_language_is_not_mistaken_for_backend_notes(self) -> None:
        self.payload["slides"][0]["title"] = "梳理木版年画的制作过程"
        self.payload["slides"][0]["teacher_notes"] = (
            "检查结论是否有诗句依据，再用自己的话修订。"
        )
        self._write_inputs()

        report = self.export()

        self.assertEqual(report["artifact"]["slide_count"], 2)

    def test_missing_font_and_unapproved_source_fail_closed(self) -> None:
        with self.assertRaisesRegex(MODULE.ExportError, r"Source Han Serif SC.*安装"):
            MODULE.export_text_pptx(
                self.source,
                self.output,
                approval_record=self.approval,
                font_checker=lambda _font: False,
            )
        approval = json.loads(self.approval.read_text(encoding="utf-8"))
        approval["inputs"][0]["sha256"] = "f" * 64
        self.approval.write_text(json.dumps(approval), encoding="utf-8")
        with self.assertRaisesRegex(MODULE.ExportError, "批准记录.*哈希"):
            self.export()


if __name__ == "__main__":
    unittest.main()
