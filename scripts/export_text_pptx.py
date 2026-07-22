#!/usr/bin/env python3
"""把已批准的逐页内容确定性导出为可直接授课的白底 PPTX。"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any, Callable, Sequence
import zipfile

try:
    import pptx
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
except ImportError:  # pragma: no cover - exercised on teacher machines without the dependency
    pptx = None


SCHEMA_VERSION = "yuwen-courseware.text-pptx-export/v1"
FONT_NAME = "Source Han Serif SC"
SLIDE_WIDTH = 12192000
SLIDE_HEIGHT = 6858000
BLACK = (0, 0, 0)
WHITE = (255, 255, 255)
FIXED_ZIP_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
BUFFER_SIZE = 1024 * 1024
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
ID_PATTERNS = {
    "slide_id": re.compile(r"^SLIDE-[0-9]{2,}$"),
    "task_id": re.compile(r"^TASK-[0-9]{2,}$"),
    "objective_ids": re.compile(r"^OBJ-[0-9]{2,}$"),
    "script_ids": re.compile(r"^SCRIPT-[0-9]{2,}$"),
    "answer_ids": re.compile(r"^ANS-[0-9]{2,}$"),
}
LABEL_NUMBER = r"(?:[一二三四五六七八九十0-9]+)?"
INTERACTION_LABEL_PATTERNS = {
    "问题": re.compile(
        rf"(?m)^[ \t]*(?:(?:核心|课堂|学习|探究)[ \t]*)?问题{LABEL_NUMBER}[ \t]*[：:]"
    ),
    "活动": re.compile(
        rf"(?m)^[ \t]*(?:(?:课堂|学习|学生|小组)[ \t]*)?活动{LABEL_NUMBER}[ \t]*[：:]"
    ),
    "时间": re.compile(
        rf"(?m)^[ \t]*(?:时间|用时|限时){LABEL_NUMBER}[ \t]*[：:]"
    ),
    "成果": re.compile(
        rf"(?m)^[ \t]*(?:成果|产出){LABEL_NUMBER}[ \t]*[：:]"
    ),
    "交流或提交方式": re.compile(
        rf"(?m)^[ \t]*(?:交流|汇报|展示|分享|提交)(?:方式)?{LABEL_NUMBER}[ \t]*[：:]"
    ),
}
INTERACTION_PLACEHOLDER_RE = re.compile(
    r"(?:\bTODO\b|\bTBD\b|\bPLACEHOLDER\b|待填写|待补充|待生成|"
    r"此处填写|此处讲解|后续补充|replace[-_ ]?me|replace-source-id)",
    re.IGNORECASE,
)
ACTUAL_INTERACTION_TEXT_RE = re.compile(r"[A-Za-z0-9\u3400-\u9fff]")
TEACHER_VISIBLE_FIELDS = {
    "title": "标题",
    "screen_text": "屏显文字",
    "teacher_notes": "教师备注",
}
BACKEND_TERMS = (
    "文本合同",
    "技术冻结",
    "哈希绑定",
    "占位语检测",
    "权利状态",
    "系统一致性合同",
    "文件哈希",
    "批准记录",
    "校验值",
)
BACKEND_ARTIFACT_RE = re.compile(
    r"(?<![A-Za-z0-9_])(?:schema[\s_-]*version|sha\s*[-_－]?\s*256|"
    r"manifest(?:\.json)?|approval[\s_-]*record|python[\s_-]*pptx)"
    r"(?![A-Za-z0-9_])|(?<![A-Za-z0-9])json(?![A-Za-z0-9])"
    r"[ \t]*(?:文件|格式|结构)?[ \t]*(?:校验|验证|检查)(?:报告|结果|记录|结论)?",
    re.IGNORECASE,
)
BACKEND_PROCESS_RE = re.compile(
    r"(?:^|[\r\n。；;])[ \t]*(?:"
    r"(?:后台|系统)(?:说明|备注|提示)|"
    r"(?:当前|本轮|本次)?(?:检查|校验|验证)(?:结论|报告)[ \t]*[：:]|"
    r"(?:课件|PPTX?|文件|文档)(?:的)?(?:制作|生成|导出)(?:过程|说明|记录)"
    r")[ \t]*[：:]?|"
    r"(?:制作|生成|导出)(?:过程|说明|记录)[ \t]*[：:]"
    r"[^\r\n。；;]{0,80}(?:脚本|程序|代码|软件|python|PPTX?|JSON|文件|文档|导出|生成|校验|兼容)",
    re.IGNORECASE,
)
SOFTWARE_COMPATIBILITY_RE = re.compile(
    r"(?:software[\s_-]*compatibility|compatibility[\s_-]*(?:notice|warning|report))|"
    r"(?:软件|课件|文件|文档|PPTX?|PowerPoint|WPS|Office|LibreOffice)"
    r"[^\r\n。；;]{0,24}(?:兼容(?:性)?|适配|打开(?:方式)?|版本(?:要求)?|"
    r"字体(?:替换|缺失)|错位|显示异常|播放异常|换行)|"
    r"(?:兼容(?:性)?|适配|打开方式|版本要求)"
    r"[^\r\n。；;]{0,24}(?:PPTX?|PowerPoint|WPS|Office|LibreOffice)",
    re.IGNORECASE,
)
TEACHER_VISIBLE_ID_RE = re.compile(
    r"(?<![A-Za-z0-9_-])(?:OBJ|TASK|ASM|ANS|SCRIPT|SLIDE|SRC|CLM|MAT)-[0-9]+"
    r"(?![A-Za-z0-9_-])",
    re.IGNORECASE,
)
BACKEND_STATUS_RE = re.compile(
    r"(?<![A-Za-z0-9_-])(?:status|verification[_-]status|package[_-]status|"
    r"external[_-]ppt[_-]status|draft|review|approved|invalidated|student-blank|prompt|"
    r"answer|summary|neutral|release-candidate|final-package-ready|not-requested|"
    r"design-handoff-ready|returned|content-checked|pass|failed|lesson-brief|"
    r"research-plan|evidence-dossier|stance-selection|route-freeze|teaching-design|"
    r"student-resources|teaching-script|text-contract|ppt-text|external-ppt-check|"
    r"release)(?![A-Za-z0-9_-])",
    re.IGNORECASE,
)
DECORATIVE_EMOJI_RE = re.compile(
    r"[\U0001F1E6-\U0001F1FF\U0001F300-\U0001FAFF\u2600-\u27BF\uFE0F\u20E3]"
)
ROOT_FIELDS = {
    "schema_version",
    "record_type",
    "lesson_id",
    "status",
    "generated_at",
    "text_contract_sha256",
    "ppt_text_sha256",
    "slides",
}
SLIDE_REQUIRED_FIELDS = {
    "slide_id",
    "page",
    "title",
    "screen_text",
    "teacher_notes",
    "task_id",
    "objective_ids",
    "script_ids",
    "answer_ids",
    "reveal_state",
}
SLIDE_OPTIONAL_FIELDS = {
    "material_direction",
    "information_hierarchy",
    "immutable_content",
}
WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{number}" for number in range(1, 10)),
    *(f"LPT{number}" for number in range(1, 10)),
}


class ExportError(RuntimeError):
    """可由教师或制作人直接处理的导出错误。"""


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(BUFFER_SIZE), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_json_without_duplicates(text: str, label: str) -> Any:
    def reject_duplicates(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for key, value in pairs:
            if key in result:
                raise ExportError(f"{label}含有重复字段：{key}")
            result[key] = value
        return result

    try:
        return json.loads(text, object_pairs_hook=reject_duplicates)
    except (UnicodeError, json.JSONDecodeError, RecursionError) as error:
        raise ExportError(f"{label}不是有效 UTF-8 JSON") from error


def read_json(path: Path, label: str) -> tuple[dict[str, Any], bytes]:
    try:
        raw = path.read_bytes()
        text = raw.decode("utf-8")
    except (OSError, UnicodeError) as error:
        raise ExportError(f"无法读取{label}：{path.name}") from error
    value = parse_json_without_duplicates(text, label)
    if not isinstance(value, dict):
        raise ExportError(f"{label}顶层必须是对象")
    return value, raw


def validate_portable_component(value: str, label: str) -> str:
    if not value or value in {".", ".."} or Path(value).name != value:
        raise ExportError(f"{label}必须是单个文件名，不能是路径")
    if "/" in value or "\\" in value or re.search(r'[<>:"|?*\x00-\x1f]', value):
        raise ExportError(f"{label}含有跨平台不兼容字符")
    if value.endswith((" ", ".")):
        raise ExportError(f"{label}不能以空格或句点结尾")
    if value.split(".", 1)[0].rstrip(" .").upper() in WINDOWS_RESERVED_NAMES:
        raise ExportError(f"{label}是 Windows 保留名称")
    return value


def require_text(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ExportError(f"{label}必须是非空文字")
    return value


def validate_structured_ppt(value: dict[str, Any]) -> list[dict[str, Any]]:
    if set(value) != ROOT_FIELDS:
        raise ExportError("PPT 结构化 JSON 顶层字段与 schema 不一致")
    if value["schema_version"] != "1.0" or value["record_type"] != "ppt-text":
        raise ExportError("PPT 结构化 JSON 的版本或记录类型无效")
    lesson_id = require_text(value["lesson_id"], "lesson_id")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}", lesson_id):
        raise ExportError("lesson_id 不是可移植标识符")
    if value["status"] not in {"draft", "review", "approved", "invalidated"}:
        raise ExportError("PPT 结构化 JSON 的 status 无效")
    for key in ("text_contract_sha256", "ppt_text_sha256"):
        if not isinstance(value[key], str) or not SHA256_RE.fullmatch(value[key]):
            raise ExportError(f"{key} 必须是小写 SHA-256")
    slides = value["slides"]
    if not isinstance(slides, list) or not slides:
        raise ExportError("PPT 结构化 JSON 至少需要一页")
    for expected_page, slide in enumerate(slides, 1):
        label = f"第 {expected_page} 页"
        if not isinstance(slide, dict):
            raise ExportError(f"{label}必须是对象")
        if set(slide) - (SLIDE_REQUIRED_FIELDS | SLIDE_OPTIONAL_FIELDS):
            raise ExportError(f"{label}含有 schema 未允许字段")
        if not SLIDE_REQUIRED_FIELDS.issubset(slide):
            raise ExportError(f"{label}缺少必填字段")
        if slide["page"] != expected_page:
            raise ExportError(f"{label}页码必须从 1 连续递增")
        require_text(slide["title"], f"{label}标题")
        require_text(slide["screen_text"], f"{label}屏显文字")
        require_text(slide["teacher_notes"], f"{label}教师备注")
        for field in ("slide_id", "task_id"):
            if not isinstance(slide[field], str) or not ID_PATTERNS[field].fullmatch(slide[field]):
                raise ExportError(f"{label}{field}格式无效")
        for field in ("objective_ids", "script_ids", "answer_ids"):
            identifiers = slide[field]
            if not isinstance(identifiers, list) or (
                field != "answer_ids" and not identifiers
            ):
                raise ExportError(f"{label}{field}必须是有效列表")
            if len(identifiers) != len(set(identifiers)) or any(
                not isinstance(identifier, str)
                or not ID_PATTERNS[field].fullmatch(identifier)
                for identifier in identifiers
            ):
                raise ExportError(f"{label}{field}含无效或重复编号")
        if slide["reveal_state"] not in {
            "student-blank",
            "prompt",
            "answer",
            "summary",
            "neutral",
        }:
            raise ExportError(f"{label}reveal_state 无效")
    return slides


def validate_approval(
    approval: dict[str, Any],
    *,
    approval_path: Path,
    source_path: Path,
    source_hash: str,
    source_size: int,
    lesson_id: str,
) -> str:
    if (
        approval.get("schema_version") != "1.0"
        or approval.get("record_type") != "user-stage-approval"
        or approval.get("stage") != "ppt-text"
        or approval.get("next_stage") != "release"
        or approval.get("lesson_id") != lesson_id
    ):
        raise ExportError("批准记录不是该课例的 PPT 文字稿批准记录")
    approval_id = require_text(approval.get("approval_id"), "approval_id")
    inputs = approval.get("inputs")
    if not isinstance(inputs, list):
        raise ExportError("批准记录缺少 inputs")
    matches = [item for item in inputs if isinstance(item, dict) and item.get("role") == "ppt-structured"]
    if len(matches) != 1:
        raise ExportError("批准记录必须且只能绑定一个 ppt-structured 输入")
    record = matches[0]
    if record.get("path") != source_path.name:
        raise ExportError("批准记录绑定的结构化 JSON 文件名不一致")
    if record.get("sha256") != source_hash or record.get("size") != source_size:
        raise ExportError("批准记录与结构化 JSON 的哈希或大小不一致")
    validate_portable_component(approval_path.name, "批准记录文件名")
    return approval_id


def _font_names_from_fc_list() -> set[str]:
    executable = shutil.which("fc-list")
    if executable is None:
        return set()
    try:
        completed = subprocess.run(
            [executable, ":", "family"],
            check=False,
            capture_output=True,
            text=True,
            timeout=20,
        )
    except (OSError, subprocess.TimeoutExpired):
        return set()
    if completed.returncode != 0:
        return set()
    return {
        family.strip()
        for line in completed.stdout.splitlines()
        for family in line.split(",")
        if family.strip()
    }


def font_is_installed(font_name: str = FONT_NAME) -> bool:
    if font_name in _font_names_from_fc_list():
        return True
    font_tokens = ("sourcehanserifsc", "source-han-serif-sc")
    roots = [Path("/Library/Fonts"), Path("/System/Library/Fonts")]
    user_fonts = Path.home() / "Library" / "Fonts"
    roots.append(user_fonts)
    for root in roots:
        if not root.is_dir():
            continue
        try:
            for path in root.rglob("*"):
                compact = path.name.casefold().replace(" ", "")
                if path.is_file() and any(token.replace("-", "") in compact.replace("-", "") for token in font_tokens):
                    return True
        except OSError:
            continue
    if os.name == "nt":  # pragma: no cover - exercised only on Windows
        try:
            import winreg

            for hive, key_name in (
                (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
                (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
            ):
                try:
                    with winreg.OpenKey(hive, key_name) as key:
                        index = 0
                        while True:
                            name, _, _ = winreg.EnumValue(key, index)
                            if font_name.casefold() in name.casefold():
                                return True
                            index += 1
                except (FileNotFoundError, OSError):
                    continue
        except ImportError:
            pass
    return False


def ensure_runtime(font_checker: Callable[[str], bool]) -> None:
    if pptx is None:
        raise ExportError(
            "缺少 python-pptx。请先运行 `<PYTHON> -m pip install python-pptx`，再重新导出。"
        )
    if not font_checker(FONT_NAME):
        raise ExportError(
            "未检测到 Source Han Serif SC（思源宋体）。请从 Adobe Source Han Serif 官方发布页"
            "安装 SC 字体并重启当前智能体后再导出；不得用其他字体静默替代："
            "https://github.com/adobe-fonts/source-han-serif/releases"
        )


def _character_units(text: str) -> float:
    units = 0.0
    for character in text:
        if character.isspace():
            units += 0.45
        elif ord(character) < 128:
            units += 0.58
        elif character in "，。！？；：、（）《》〈〉“”‘’｜—…":
            units += 0.8
        else:
            units += 1.0
    return units


def is_section_line(line: str, index: int) -> bool:
    stripped = line.strip()
    if not stripped or stripped.startswith(tuple("①②③④⑤⑥⑦⑧⑨0123456789")):
        return False
    if stripped.endswith(("：", ":")) and len(stripped) <= 24:
        return True
    return index == 0 and len(stripped) <= 18 and not re.search(r"[。！？]$", stripped)


def body_line_sizes(screen_text: str, cover: bool) -> list[tuple[str, int]]:
    result: list[tuple[str, int]] = []
    non_empty_index = 0
    for line in screen_text.splitlines():
        if not line.strip():
            result.append(("", 28))
            continue
        if cover:
            size = 36
        else:
            size = 32 if is_section_line(line, non_empty_index) else 28
        result.append((line, size))
        non_empty_index += 1
    return result


def overflow_lines(screen_text: str, *, cover: bool) -> int:
    width_inches = 11.55
    height_points = (4.55 if cover else 5.35) * 72
    used_points = 0.0
    for line, font_size in body_line_sizes(screen_text, cover):
        if not line:
            used_points += 16
            continue
        capacity = max(1.0, width_inches * 72 / font_size)
        wrapped = max(1, math.ceil(_character_units(line) / capacity))
        used_points += wrapped * font_size * 1.18
    if used_points <= height_points:
        return 0
    return max(1, math.ceil((used_points - height_points) / (28 * 1.18)))


def validate_layout_capacity(slides: list[dict[str, Any]]) -> None:
    for slide in slides:
        page = slide["page"]
        cover = page == 1
        title_capacity = 12 if cover else 24
        title_lines = max(1, math.ceil(_character_units(slide["title"]) / title_capacity))
        allowed_title_lines = 2 if cover else 1
        if title_lines > allowed_title_lines:
            raise ExportError(
                f"第 {page} 页标题超出 {title_lines - allowed_title_lines} 行；请回到内容阶段减字，不会自动缩小字号。"
            )
        excess = overflow_lines(slide["screen_text"], cover=cover)
        if excess:
            raise ExportError(
                f"第 {page} 页屏显文字预计超出 {excess} 行；请回到内容阶段减字，不会自动缩小字号。"
            )


def validate_student_visible_interactions(slides: list[dict[str, Any]]) -> None:
    """确保每个课堂任务都能只靠学生可见页面完整实施。"""

    task_groups: dict[str, list[dict[str, Any]]] = {}
    for slide in slides:
        task_groups.setdefault(slide["task_id"], []).append(slide)

    for group in task_groups.values():
        visible_text = "\n".join(slide["screen_text"] for slide in group)
        first_page = group[0]["page"]
        incomplete = [
            label
            for label, pattern in INTERACTION_LABEL_PATTERNS.items()
            if not interaction_label_has_actual_content(visible_text, pattern)
        ]
        if incomplete:
            incomplete_text = "、".join(incomplete)
            raise ExportError(
                f"第 {first_page} 页所在的课堂任务缺少学生可见的{incomplete_text}实际内容；"
                "请在标签后的同一行或下一行写完整，不能留空、写占位语，"
                "也不能只写在教师备注或学案里。"
            )


def interaction_label_has_actual_content(text: str, pattern: re.Pattern[str]) -> bool:
    """读取标签后的同一行或后续首个内容行，不把下一标签当作内容。"""

    for match in pattern.finditer(text):
        for raw_line in text[match.end():].splitlines():
            candidate = raw_line.strip()
            if not candidate:
                continue
            if any(
                other_pattern.match(raw_line)
                for other_pattern in INTERACTION_LABEL_PATTERNS.values()
            ):
                break
            if INTERACTION_PLACEHOLDER_RE.search(candidate):
                break
            if ACTUAL_INTERACTION_TEXT_RE.search(candidate):
                return True
    return False


def teacher_facing_interference(value: str) -> str | None:
    if (
        any(term in value for term in BACKEND_TERMS)
        or BACKEND_ARTIFACT_RE.search(value)
        or BACKEND_PROCESS_RE.search(value)
    ):
        return "后台制作或检查说明"
    if SOFTWARE_COMPATIBILITY_RE.search(value):
        return "软件兼容提醒"
    if TEACHER_VISIBLE_ID_RE.search(value):
        return "内部编号"
    if BACKEND_STATUS_RE.search(value):
        return "英文状态词"
    if DECORATIVE_EMOJI_RE.search(value):
        return "装饰性图标"
    return None


def validate_teacher_facing_text(slides: list[dict[str, Any]]) -> None:
    for slide in slides:
        for field, field_label in TEACHER_VISIBLE_FIELDS.items():
            interference = teacher_facing_interference(slide[field])
            if interference is not None:
                raise ExportError(
                    f"第 {slide['page']} 页{field_label}含有{interference}；"
                    "请只保留课堂教学需要的自然中文。"
                )


def _set_east_asian_font(run: Any, font_name: str) -> None:
    properties = run._r.get_or_add_rPr()
    east_asian = properties.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        properties.append(east_asian)
    east_asian.set("typeface", font_name)
    properties.set("lang", "zh-CN")


def style_run(run: Any, size: int, *, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    _set_east_asian_font(run, FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*BLACK)


def add_text_box(
    slide: Any,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    lines: list[tuple[str, int]],
    bold: bool,
    align: Any,
    vertical_anchor: Any,
) -> Any:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = vertical_anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, (text, size) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = 1.08
        paragraph.space_after = Pt(5 if text else 0)
        if text:
            run = paragraph.add_run()
            run.text = text
            style_run(run, size, bold=bold)
    return shape


def render_presentation(slides: list[dict[str, Any]], output_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    properties = presentation.core_properties
    properties.title = slides[0]["title"]
    properties.subject = "高中语文白底课堂课件"
    properties.author = "语文优课"
    properties.last_modified_by = "语文优课"
    properties.created = datetime(2000, 1, 1)
    properties.modified = datetime(2000, 1, 1)
    properties.revision = 1
    blank_layout = presentation.slide_layouts[6]
    for source_slide in slides:
        cover = source_slide["page"] == 1
        slide = presentation.slides.add_slide(blank_layout)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(*WHITE)
        add_text_box(
            slide,
            left=0.85,
            top=0.45 if cover else 0.38,
            width=11.63,
            height=1.65 if cover else 0.78,
            lines=[(source_slide["title"], 80 if cover else 40)],
            bold=True,
            align=PP_ALIGN.CENTER if cover else PP_ALIGN.LEFT,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text_box(
            slide,
            left=0.9,
            top=2.18 if cover else 1.35,
            width=11.55,
            height=4.55 if cover else 5.35,
            lines=body_line_sizes(source_slide["screen_text"], cover),
            bold=False,
            align=PP_ALIGN.CENTER if cover else PP_ALIGN.LEFT,
            vertical_anchor=MSO_ANCHOR.TOP,
        )
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = source_slide["teacher_notes"]
        for paragraph in notes_frame.paragraphs:
            for run in paragraph.runs:
                style_run(run, 12)
    presentation.save(output_path)


def normalize_pptx_archive(path: Path) -> None:
    normalized = path.with_name(path.stem + ".normalized.pptx")
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = 0o644 << 16
                target.writestr(info, source.read(name))
        normalized.replace(path)
    except (OSError, zipfile.BadZipFile) as error:
        try:
            normalized.unlink()
        except FileNotFoundError:
            pass
        raise ExportError("无法规范化 PPTX，未生成交付物") from error


def validate_pptx(path: Path, expected_slides: int) -> None:
    if not path.is_file() or path.stat().st_size == 0 or not zipfile.is_zipfile(path):
        raise ExportError("未生成有效 PPTX")
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if not {"[Content_Types].xml", "ppt/presentation.xml"}.issubset(names):
                raise ExportError("PPTX 缺少必要 OOXML 部件")
            slide_count = sum(
                1
                for name in names
                if re.fullmatch(r"ppt/slides/slide[0-9]+\.xml", name)
            )
    except (OSError, zipfile.BadZipFile) as error:
        raise ExportError("PPTX 不是有效的 OOXML 压缩包") from error
    if slide_count != expected_slides:
        raise ExportError("PPTX 页数与结构化 JSON 不一致")


def isoformat_utc(value: datetime) -> str:
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def export_text_pptx(
    source: Path,
    output_dir: Path,
    *,
    approval_record: Path,
    name: str | None = None,
    expected_source_sha256: str | None = None,
    font_checker: Callable[[str], bool] = font_is_installed,
    created_at: datetime | None = None,
) -> dict[str, Any]:
    ensure_runtime(font_checker)
    source = source.resolve()
    approval_record = approval_record.resolve()
    if source.suffix.lower() != ".json" or not source.is_file():
        raise ExportError("输入必须是存在的 PPT 结构化 JSON 文件")
    if not approval_record.is_file():
        raise ExportError("必须提供存在的 PPT 文字稿批准记录")
    validate_portable_component(source.name, "PPT 结构化 JSON 文件名")
    output_name = validate_portable_component(name or source.stem, "输出名称")
    payload, source_bytes = read_json(source, "PPT 结构化 JSON")
    slides = validate_structured_ppt(payload)
    source_hash = sha256_bytes(source_bytes)
    if expected_source_sha256 is not None and (
        not SHA256_RE.fullmatch(expected_source_sha256) or expected_source_sha256 != source_hash
    ):
        raise ExportError("结构化 JSON 的冻结哈希与 --expect-source-sha256 不一致")
    approval, approval_bytes = read_json(approval_record, "批准记录")
    approval_id = validate_approval(
        approval,
        approval_path=approval_record,
        source_path=source,
        source_hash=source_hash,
        source_size=len(source_bytes),
        lesson_id=payload["lesson_id"],
    )
    validate_teacher_facing_text(slides)
    validate_student_visible_interactions(slides)
    validate_layout_capacity(slides)
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    if not output_dir.is_dir():
        raise ExportError("输出目录无效")
    final_bundle = output_dir / f"{output_name}.export"
    if final_bundle.exists() or final_bundle.is_symlink():
        raise ExportError(f"拒绝覆盖已有导出目录：{final_bundle.name}")
    lock = output_dir / f".{output_name}.pptx.lock"
    try:
        descriptor = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError as error:
        raise ExportError("同名 PPTX 导出正在进行") from error
    staging = Path(tempfile.mkdtemp(prefix=f".{output_name}.pptx-", dir=output_dir))
    try:
        pptx_path = staging / f"{output_name}.pptx"
        approval_snapshot = staging / f"{output_name}.批准记录.json"
        approval_snapshot.write_bytes(approval_bytes)
        render_presentation(slides, pptx_path)
        normalize_pptx_archive(pptx_path)
        validate_pptx(pptx_path, len(slides))
        artifact_hash = sha256_file(pptx_path)
        timestamp = created_at or datetime.now(timezone.utc)
        report = {
            "schema_version": SCHEMA_VERSION,
            "ok": True,
            "created_at": isoformat_utc(timestamp),
            "path_base": "export-bundle-directory",
            "bundle": {"path": final_bundle.name, "path_kind": "output-relative"},
            "commit": {"mode": "atomic-directory-rename", "complete": True},
            "source": {
                "path": source.name,
                "path_kind": "basename-only",
                "sha256": source_hash,
                "size": len(source_bytes),
                "frozen_during_export": True,
            },
            "approval": {
                "path": approval_snapshot.name,
                "path_kind": "report-relative",
                "sha256": sha256_bytes(approval_bytes),
                "size": len(approval_bytes),
                "approval_id": approval_id,
            },
            "artifact": {
                "path": pptx_path.name,
                "sha256": artifact_hash,
                "size": pptx_path.stat().st_size,
                "slide_count": len(slides),
                "slide_width_emu": SLIDE_WIDTH,
                "slide_height_emu": SLIDE_HEIGHT,
                "font": FONT_NAME,
            },
            "tools": {
                "python": sys.version.split()[0],
                "python_pptx": pptx.__version__,
            },
            "lineage": {
                "description": "已批准的逐页内容生成可直接授课的白底 PPTX",
                "layout_policy": "16:9-white-black-no-autoshrink",
            },
        }
        report_path = staging / f"{output_name}.export-report.json"
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        final_bundle.parent.mkdir(parents=True, exist_ok=True)
        staging.replace(final_bundle)
        return report
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    finally:
        os.close(descriptor)
        try:
            lock.unlink()
        except FileNotFoundError:
            pass


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="把已批准的逐页内容导出为可直接授课的白底 PPTX")
    parser.add_argument("source", type=Path, help="已批准的 ppt-text-structured.json")
    parser.add_argument("output_dir", type=Path, help="导出目录")
    parser.add_argument("--approval-record", required=True, type=Path, help="PPT 文字稿批准记录 JSON")
    parser.add_argument("--name", help="输出文件名（不含扩展名）")
    parser.add_argument("--expect-source-sha256", help="冻结的结构化 JSON SHA-256")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        report = export_text_pptx(
            args.source,
            args.output_dir,
            approval_record=args.approval_record,
            name=args.name,
            expected_source_sha256=args.expect_source_sha256,
        )
    except ExportError as error:
        print(f"导出失败：{error}", file=sys.stderr)
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
